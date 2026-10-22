# -*- coding: utf-8 -*-
"""提取 docx 纯文本，供资料上传后解析使用。"""
import argparse
import json
import sys

from docx import Document


def main() -> None:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")
        except Exception:
            pass

    parser = argparse.ArgumentParser()
    parser.add_argument("--file", required=True)
    args = parser.parse_args()

    doc = Document(args.file)
    parts = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    text = "\n".join(parts)
    print("JSWJ_TEXT " + json.dumps({"text": text, "chars": len(text)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
